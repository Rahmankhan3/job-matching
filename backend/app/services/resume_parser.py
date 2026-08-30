"""
Resume parsing service.
Extracts structured candidate data from PDF / DOCX resumes and returns
a validated ParsedResume model. Uses spaCy for NER and the skill normalizer
for consistent skill identification with provenance tracking.
"""
import os
import re
import logging
from typing import Optional, List
from datetime import datetime

import pdfplumber
from docx import Document

from app.config import settings
from app.services.model_loader import get_nlp
from app.services.skill_normalizer import (
    extract_skills_from_text,
    extract_skills_with_evidence,
    normalize_skills,
)
from app.models.parsed_resume import (
    ParsedResume, PersonalInfo, SkillEvidence,
    ExperienceEntry, ProjectEntry, EducationEntry, CertificationEntry,
)

logger = logging.getLogger(__name__)

# ── Regex Patterns ─────────────────────────────────────────────────────────────
EMAIL_RE = re.compile(r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}')
PHONE_RE = re.compile(r'(?:\+?\d{1,3}[\s.-]?)?\(?\d{2,4}\)?[\s.-]?\d{3,5}[\s.-]?\d{3,5}')
LINKEDIN_RE = re.compile(r'linkedin\.com/in/[\w-]+', re.IGNORECASE)
GITHUB_RE = re.compile(r'github\.com/[\w-]+', re.IGNORECASE)
YEAR_RE = re.compile(r'\b(19|20)\d{2}\b')
DURATION_RE = re.compile(
    r'(\w+[\s]*\d{4})\s*[-–—to]+\s*(\w+[\s]*\d{4}|present|current|till date)',
    re.IGNORECASE,
)

# Month name → number for duration estimation
MONTH_MAP = {
    "jan": 1, "january": 1, "feb": 2, "february": 2,
    "mar": 3, "march": 3, "apr": 4, "april": 4,
    "may": 5, "jun": 6, "june": 6, "jul": 7, "july": 7,
    "aug": 8, "august": 8, "sep": 9, "sept": 9, "september": 9,
    "oct": 10, "october": 10, "nov": 11, "november": 11,
    "dec": 12, "december": 12,
}

# Section header patterns — resilient to common naming variations
SECTION_HEADERS = {
    'experience': re.compile(
        r'\b(experience|work\s*history|employment|professional\s*experience'
        r'|work\s*experience|career\s*history)\b', re.I
    ),
    'education': re.compile(
        r'\b(education|academic|qualification|degree|educational\s*background)\b', re.I
    ),
    'projects': re.compile(
        r'\b(project|personal\s*project|academic\s*project|key\s*project'
        r'|notable\s*project|side\s*project)\b', re.I
    ),
    'skills': re.compile(
        r'\b(skill|technical\s*skill|competenc|technolog|proficiency'
        r'|core\s*competenc|expertise|tools?\s*(?:and|&)\s*technologies)\b', re.I
    ),
    'certifications': re.compile(
        r'\b(certification|certificate|accreditation|credential'
        r'|professional\s*certification|licenses?\s*(?:and|&)\s*certifications?)\b', re.I
    ),
    'summary': re.compile(
        r'\b(summary|about\s*me|profile|objective|professional\s*summary'
        r'|career\s*objective|career\s*summary|overview)\b', re.I
    ),
}

# Education degree patterns — ordered from highest to lowest
DEGREE_PATTERNS = [
    re.compile(r'\b(Ph\.?D|Doctor of Philosophy)\b', re.I),
    re.compile(r'\b(M\.?Tech|M\.?S\.?|Master|MBA|M\.?E\.?|M\.?Sc|M\.?A\.?|M\.?Com)\b', re.I),
    re.compile(r'\b(B\.?Tech|B\.?E\.?|B\.?S\.?|Bachelor|BCA|BBA|B\.?Sc|B\.?A\.?|B\.?Com)\b', re.I),
    re.compile(r'\b(Diploma|Associate|Higher Secondary|HSC|SSC|12th|10th)\b', re.I),
]

# Common education fields
FIELD_PATTERNS = re.compile(
    r'(Computer Science|Information Technology|Electronics|Electrical|Mechanical'
    r'|Civil|Data Science|Software|Mathematics|Physics|Chemistry|Commerce'
    r'|Business|Management|Arts|Science|Engineering)', re.I
)

# Certification patterns
CERT_ISSUERS = re.compile(
    r'(AWS|Google|Microsoft|Azure|Coursera|Udemy|Oracle|Cisco|CompTIA'
    r'|IBM|Meta|HackerRank|Salesforce|PMI|Scrum\.org)', re.I
)


# ── Text Extraction ───────────────────────────────────────────────────────────

def _extract_text_from_pdf(file_path: str) -> str:
    text_parts = []
    try:
        with pdfplumber.open(file_path) as pdf:
            if not pdf.pages:
                logger.warning(f"PDF has no pages: {file_path}")
                return ""
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
    except Exception as e:
        logger.error(f"PDF extraction failed for {file_path}: {e}")

    extracted = "\n".join(text_parts).strip()
    
    # If text is empty or minimal (< 30 chars), the PDF may be scanned / image-rendered.
    # Attempt OCR fallback if OCR libraries are installed.
    if len(extracted) < 30:
        try:
            import pypdfium2 as pdfium
            from rapidocr_onnxruntime import RapidOCR
            ocr = RapidOCR()
            pdf = pdfium.PdfDocument(file_path)
            try:
                ocr_text_parts = []
                for page_idx in range(len(pdf)):
                    page = pdf[page_idx]
                    image = page.render(scale=2.0).to_numpy()
                    ocr_result, _ = ocr(image)
                    if ocr_result:
                        page_ocr_text = "\n".join([line[1] for line in ocr_result])
                        ocr_text_parts.append(page_ocr_text)
                if ocr_text_parts:
                    extracted = "\n".join(ocr_text_parts).strip()
                    logger.info(f"Successfully extracted {len(extracted)} chars via OCR for {os.path.basename(file_path)}")
            finally:
                pdf.close()
        except Exception as ocr_err:
            logger.debug(f"OCR fallback unavailable or failed for {file_path}: {ocr_err}")

    return extracted


def _extract_text_from_docx(file_path: str) -> str:
    text_parts = []
    try:
        doc = Document(file_path)
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)
    except Exception as e:
        logger.error(f"DOCX extraction failed for {file_path}: {e}")
    return "\n".join(text_parts)


def extract_text(file_path: str) -> str:
    """Extract raw text from a resume file (PDF or DOCX)."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return _extract_text_from_pdf(file_path)
    elif ext in (".docx", ".doc"):
        return _extract_text_from_docx(file_path)
    return ""


# ── Text Cleaning ─────────────────────────────────────────────────────────────

def _clean_text(text: str) -> str:
    """Basic text cleaning: normalize whitespace, remove control characters."""
    # Remove null bytes and control chars (except newlines/tabs)
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', text)
    # Normalize multiple spaces (but keep newlines for section splitting)
    text = re.sub(r'[^\S\n]+', ' ', text)
    # Remove excessive blank lines
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()


# ── Section Splitting ─────────────────────────────────────────────────────────

def _split_sections(text: str) -> dict:
    """Split resume text into rough sections based on header patterns."""
    lines = text.split("\n")
    sections = {
        'header': [], 'experience': [], 'education': [],
        'projects': [], 'skills': [], 'certifications': [],
        'summary': [], 'other': [],
    }
    current = 'header'

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        matched_section = None
        # Only match short lines that look like headers
        if len(stripped) < 60:
            for sec_name, pattern in SECTION_HEADERS.items():
                if pattern.search(stripped):
                    matched_section = sec_name
                    break

        if matched_section:
            current = matched_section
        else:
            sections[current].append(stripped)

    return sections


# ── Extraction Functions ──────────────────────────────────────────────────────

def _extract_personal_info(text: str, sections: dict) -> PersonalInfo:
    """Extract name, email, phone, location, links from the header area."""
    nlp = get_nlp()
    header_text = "\n".join(sections.get('header', [])[:10])
    doc = nlp(header_text) if header_text else nlp(text[:500])

    # Name — first PERSON entity in the header
    name = None
    for ent in doc.ents:
        if ent.label_ == "PERSON":
            name = ent.text.strip()
            break

    # Fallback: first line that looks like a name
    if not name and sections.get('header'):
        first_line = sections['header'][0]
        if len(first_line.split()) <= 5 and not EMAIL_RE.search(first_line):
            name = first_line.strip()

    emails = EMAIL_RE.findall(text)
    phones = PHONE_RE.findall(text)
    linkedin_matches = LINKEDIN_RE.findall(text)
    github_matches = GITHUB_RE.findall(text)

    # Location — GPE entities near the header
    location = None
    for ent in doc.ents:
        if ent.label_ in ("GPE", "LOC"):
            location = ent.text.strip()
            break

    return PersonalInfo(
        name=name,
        email=emails[0] if emails else None,
        phone=phones[0].strip() if phones else None,
        location=location,
        linkedin=f"https://{linkedin_matches[0]}" if linkedin_matches else None,
        github=f"https://{github_matches[0]}" if github_matches else None,
    )


def _estimate_duration_months(duration_str: str) -> Optional[int]:
    """Attempt to estimate duration in months from a date range string."""
    if not duration_str:
        return None

    match = DURATION_RE.search(duration_str)
    if not match:
        return None

    start_str = match.group(1).strip().lower()
    end_str = match.group(2).strip().lower()

    def parse_month_year(s):
        parts = s.split()
        month, year = 1, None
        for p in parts:
            p_clean = p.strip().lower()
            if p_clean in MONTH_MAP:
                month = MONTH_MAP[p_clean]
            elif re.match(r'^\d{4}$', p_clean):
                year = int(p_clean)
        return month, year

    start_month, start_year = parse_month_year(start_str)
    if end_str in ("present", "current", "till date"):
        end_month = datetime.now().month
        end_year = datetime.now().year
    else:
        end_month, end_year = parse_month_year(end_str)

    if start_year and end_year:
        months = (end_year - start_year) * 12 + (end_month - start_month)
        return max(months, 0)
    return None


def _extract_experience(sections: dict) -> List[ExperienceEntry]:
    """Extract work experience entries with technology detection."""
    exp_lines = sections.get('experience', [])
    if not exp_lines:
        return []

    entries = []
    current = None

    for line in exp_lines:
        duration_match = DURATION_RE.search(line)

        is_entry_start = (
            len(line.split()) <= 10
            and not line.startswith(("•", "-", "–", "*", "◦"))
        ) or duration_match

        if is_entry_start and not line.startswith(("•", "-", "–", "*")):
            if current:
                entries.append(current)

            duration_str = duration_match.group(0) if duration_match else None
            current = {
                "role": line.strip(),
                "company": "",
                "duration": duration_str,
                "duration_months": _estimate_duration_months(duration_str),
                "description": "",
                "technologies": [],
            }
        elif current:
            current["description"] += " " + line.strip()
            skills = extract_skills_from_text(line)
            current["technologies"].extend(skills)

    if current:
        entries.append(current)

    result = []
    for e in entries:
        techs = sorted(set(e["technologies"]))
        result.append(ExperienceEntry(
            role=e["role"],
            company=e["company"] or None,
            duration=e["duration"],
            duration_months=e["duration_months"],
            description=e["description"].strip(),
            technologies=techs,
        ))

    return result[:10]


def _extract_projects(sections: dict) -> List[ProjectEntry]:
    """Extract project entries with technology detection."""
    project_lines = sections.get('projects', [])
    if not project_lines:
        return []

    projects = []
    current = None

    for line in project_lines:
        is_title = (
            len(line.split()) <= 12
            and not line.startswith(("•", "-", "–", "*", "◦"))
            and not DURATION_RE.search(line)
        )

        if is_title and len(line) > 3:
            if current:
                projects.append(current)
            current = {"title": line.strip(), "description": "", "technologies": []}
        elif current:
            current["description"] += " " + line.strip()
            skills = extract_skills_from_text(line)
            current["technologies"].extend(skills)

    if current:
        projects.append(current)

    result = []
    for p in projects:
        techs = sorted(set(p["technologies"]))
        result.append(ProjectEntry(
            title=p["title"],
            description=p["description"].strip(),
            technologies=techs,
        ))

    return result[:10]


def _extract_education(sections: dict) -> List[EducationEntry]:
    """Extract education entries with degree, field, and institution detection."""
    edu_lines = sections.get('education', [])
    if not edu_lines:
        return []

    entries = []
    current = None

    for line in edu_lines:
        degree_found = None
        for pattern in DEGREE_PATTERNS:
            m = pattern.search(line)
            if m:
                degree_found = m.group(0)
                break

        years = YEAR_RE.findall(line)

        # Detect field of study
        field_match = FIELD_PATTERNS.search(line)

        if degree_found:
            if current:
                entries.append(current)

            institution = line.replace(degree_found, "").strip(" -–—|,")
            if field_match:
                institution = institution.replace(field_match.group(0), "").strip(" -–—|,")

            current = {
                "degree": degree_found,
                "field": field_match.group(0) if field_match else None,
                "institution": institution if len(institution) > 2 else None,
                "graduation_year": int(max(years)) if years else None,
            }
        elif current:
            if not current.get("institution") or len(current["institution"]) < 3:
                current["institution"] = line.strip()
            if not current.get("field") and field_match:
                current["field"] = field_match.group(0)
            if years and not current.get("graduation_year"):
                current["graduation_year"] = int(max(years))

    if current:
        entries.append(current)

    return [EducationEntry(**e) for e in entries[:5]]


def _extract_certifications(sections: dict) -> List[CertificationEntry]:
    """Extract certification entries."""
    cert_lines = sections.get('certifications', [])
    if not cert_lines:
        return []

    certs = []
    for line in cert_lines:
        if len(line.strip()) < 5:
            continue

        issuer_match = CERT_ISSUERS.search(line)
        years = YEAR_RE.findall(line)

        name = line.strip()
        if issuer_match:
            # Try to clean the cert name by removing the issuer mention
            name = name.replace(issuer_match.group(0), "").strip(" -–—|,·")
            if not name:
                name = line.strip()

        certs.append(CertificationEntry(
            name=name,
            issuer=issuer_match.group(0) if issuer_match else None,
            year=int(max(years)) if years else None,
        ))

    return certs[:10]


def _extract_summary(sections: dict) -> Optional[str]:
    """Extract the summary/about/objective section."""
    summary_lines = sections.get('summary', [])
    if not summary_lines:
        return None
    return " ".join(summary_lines).strip() or None


def _build_skill_evidence(
    text: str,
    sections: dict,
    experience: List[ExperienceEntry],
    projects: List[ProjectEntry],
) -> tuple:
    """
    Build the unified skill list with provenance tracking.
    Returns (deduplicated_skills, skill_evidence_list).
    """
    evidence_map: dict = {}  # canonical_skill_lower -> SkillEvidence

    def _add_evidence(skill: str, source: str):
        key = skill.lower()
        if key not in evidence_map:
            evidence_map[key] = SkillEvidence(skill=skill, sources=[source])
        elif source not in evidence_map[key].sources:
            evidence_map[key].sources.append(source)

    # Skills from the skills section
    skills_text = " ".join(sections.get('skills', []))
    for s in extract_skills_from_text(skills_text):
        _add_evidence(s, "skills_section")

    # Skills from experience
    for i, exp in enumerate(experience):
        for tech in exp.technologies:
            _add_evidence(tech, f"experience:{i}")

    # Skills from projects
    for i, proj in enumerate(projects):
        for tech in proj.technologies:
            _add_evidence(tech, f"project:{i}")

    # Skills from full text (catches anything missed by section-specific extraction)
    for s in extract_skills_from_text(text):
        key = s.lower()
        if key not in evidence_map:
            _add_evidence(s, "full_text")

    all_skills = sorted(ev.skill for ev in evidence_map.values())
    evidence_list = sorted(evidence_map.values(), key=lambda e: e.skill)

    return all_skills, evidence_list


# ── Public API ────────────────────────────────────────────────────────────────

def parse_resume(file_path: str) -> Optional[ParsedResume]:
    """
    Parse a resume file and return a structured ParsedResume.
    Returns None if the file doesn't exist or contains no extractable text.
    """
    if not os.path.isfile(file_path):
        logger.warning(f"Resume file not found: {file_path}")
        return None

    raw_text = extract_text(file_path)
    if not raw_text or len(raw_text.strip()) < 20:
        logger.warning(f"Resume has insufficient text content: {file_path}")
        return None

    cleaned_text = _clean_text(raw_text)
    sections = _split_sections(cleaned_text)

    personal_info = _extract_personal_info(cleaned_text, sections)
    experience = _extract_experience(sections)
    projects = _extract_projects(sections)
    education = _extract_education(sections)
    certifications = _extract_certifications(sections)
    summary = _extract_summary(sections)

    skills, skill_evidence = _build_skill_evidence(
        cleaned_text, sections, experience, projects
    )

    return ParsedResume(
        personal_info=personal_info,
        skills=skills,
        skill_evidence=skill_evidence,
        experience=experience,
        projects=projects,
        education=education,
        certifications=certifications,
        summary=summary,
        raw_text=cleaned_text[:10000],
        parser_version=settings.PARSER_VERSION,
        parsed_at=datetime.utcnow(),
    )


def parse_resume_to_dict(file_path: str) -> Optional[dict]:
    """Parse and return as a plain dict (for MongoDB storage)."""
    result = parse_resume(file_path)
    if result is None:
        return None
    return result.model_dump(mode="json")
